import io
import json
import re
import zipfile
from datetime import datetime
from django.utils import timezone
from django.utils.text import slugify
from notes.models import Note, Notebook


def sanitize_filename(name: str, fallback: str = "sem-titulo") -> str:
    """
    Higieniza uma string para torná-la um nome de arquivo ou pasta seguro.
    """
    if not name:
        return fallback
    clean = re.sub(r'[\\/*?:"<>|]', "", name).strip()
    return clean if clean else fallback


def build_note_markdown(note: Note) -> str:
    """
    Converte uma anotação em um arquivo Markdown completo com metadados no topo.
    """
    notebook_name = note.notebook.title if note.notebook else "Sem Caderno"
    tags_list = [t.name for t in note.tags.all()]
    tags_str = ", ".join(tags_list) if tags_list else "nenhuma"

    created_str = note.created_at.strftime("%d/%m/%Y %H:%M:%S")
    updated_str = note.updated_at.strftime("%d/%m/%Y %H:%M:%S")

    title = note.title.strip() if note.title else "Sem Título"
    content = note.content_plain or ""

    md_lines = [
        f"# {title}",
        "",
        f"> **Criado em:** {created_str}  ",
        f"> **Atualizado em:** {updated_str}  ",
        f"> **Caderno:** {notebook_name}  ",
        f"> **Tags:** {tags_str}",
        "",
        "---",
        "",
        content
    ]
    return "\n".join(md_lines)


def get_notebook_full_path(notebook: Notebook) -> str:
    """
    Retorna o caminho de pastas hierárquico caso haja subcadernos.
    Ex: 'Trabalho/Projetos 2026'
    """
    parts = []
    curr = notebook
    while curr:
        parts.insert(0, sanitize_filename(curr.title, fallback="Caderno"))
        curr = curr.parent
    return "/".join(parts)


def generate_user_backup_zip(user) -> bytes:
    """
    Gera um arquivo ZIP em memória contendo o backup integral de todas as anotações
    ativas do usuário, estruturado em pastas de cadernos e com manifest.json.
    """
    in_memory = io.BytesIO()

    with zipfile.ZipFile(in_memory, mode="w", compression=zipfile.ZIP_DEFLATED) as zf:
        # 1. Recupera anotações e cadernos ativos
        notes = (
            Note.objects.filter(user=user, deleted_at__isnull=True)
            .select_related("notebook")
            .prefetch_related("tags")
            .order_by("notebook__title", "title")
        )
        notebooks = Notebook.objects.filter(user=user)

        # 2. Gera manifest.json de controle
        manifest_data = {
            "app": "skdnote",
            "version": "1.0",
            "exported_at": timezone.now().isoformat(),
            "user_email": user.email,
            "total_notes": notes.count(),
            "total_notebooks": notebooks.count(),
            "notes_meta": [
                {
                    "id": str(n.id),
                    "title": n.title,
                    "notebook": n.notebook.title if n.notebook else None,
                    "created_at": n.created_at.isoformat(),
                    "updated_at": n.updated_at.isoformat(),
                    "is_favorite": n.is_favorite,
                    "is_pinned": n.is_pinned,
                }
                for n in notes
            ]
        }
        zf.writestr("manifest.json", json.dumps(manifest_data, indent=2, ensure_ascii=False))

        # 3. Empacota cada nota em formato Markdown estruturado por pastas
        used_paths = set()
        for note in notes:
            if note.notebook:
                folder_path = get_notebook_full_path(note.notebook)
            else:
                folder_path = "Sem Caderno"

            safe_title = sanitize_filename(note.title, fallback="Nota")
            file_path = f"{folder_path}/{safe_title}.md"

            # Trata colisões de nomes dentro da mesma pasta
            counter = 1
            while file_path in used_paths:
                file_path = f"{folder_path}/{safe_title} ({counter}).md"
                counter += 1

            used_paths.add(file_path)

            markdown_content = build_note_markdown(note)
            zf.writestr(file_path, markdown_content.encode("utf-8"))

    return in_memory.getvalue()


def parse_html_to_blocks(content_html: str, content_plain: str):
    """
    Analisa o conteúdo da anotação e divide em blocos estruturados (títulos, listas,
    citações, código, parágrafos) para renderização consistente em PDF e DOCX.
    """
    if not content_html or not content_html.strip():
        blocks = []
        for line in (content_plain or "").splitlines():
            line_str = line.strip()
            if not line_str:
                continue
            if line_str.startswith('# '):
                blocks.append(('h1', line_str[2:]))
            elif line_str.startswith('## '):
                blocks.append(('h2', line_str[3:]))
            elif line_str.startswith('### '):
                blocks.append(('h3', line_str[4:]))
            elif line_str.startswith('- [ ] ') or line_str.startswith('* [ ] '):
                blocks.append(('todo', (False, line_str[6:])))
            elif line_str.startswith('- [x] ') or line_str.startswith('* [x] '):
                blocks.append(('todo', (True, line_str[6:])))
            elif line_str.startswith('- ') or line_str.startswith('* '):
                blocks.append(('bullet', line_str[2:]))
            elif line_str in ('---', '***'):
                blocks.append(('hr', ''))
            else:
                blocks.append(('p', line_str))
        return blocks

    from html.parser import HTMLParser

    class SimpleBlockParser(HTMLParser):
        def __init__(self):
            super().__init__()
            self.blocks = []
            self.current_tag = None
            self.buffer = []
            self.current_todo_checked = False

        def handle_starttag(self, tag, attrs):
            attr_dict = dict(attrs)
            if tag in ['h1', 'h2', 'h3', 'h4', 'p', 'blockquote', 'pre', 'li']:
                self.flush_buffer()
                self.current_tag = tag
            elif tag == 'div' and 'upnote-todo-item' in attr_dict.get('class', ''):
                self.flush_buffer()
                self.current_tag = 'todo'
                self.current_todo_checked = 'checked' in attr_dict.get('class', '')
            elif tag == 'input' and attr_dict.get('type') == 'checkbox':
                if 'checked' in attr_dict or attr_dict.get('checked') == 'checked':
                    self.current_todo_checked = True
            elif tag == 'hr':
                self.flush_buffer()
                self.blocks.append(('hr', ''))
            elif tag == 'br':
                self.buffer.append('\n')

        def handle_endtag(self, tag):
            if tag in ['h1', 'h2', 'h3', 'h4', 'p', 'blockquote', 'pre', 'li'] or (tag == 'div' and self.current_tag == 'todo'):
                self.flush_buffer()
                self.current_tag = None

        def handle_data(self, data):
            if self.current_tag:
                self.buffer.append(data)
            else:
                cleaned = data.strip()
                if cleaned:
                    self.blocks.append(('p', cleaned))

        def flush_buffer(self):
            text = "".join(self.buffer).strip()
            self.buffer = []
            if not text and self.current_tag != 'hr':
                return
            if self.current_tag == 'todo':
                self.blocks.append(('todo', (self.current_todo_checked, text)))
                self.current_todo_checked = False
            elif self.current_tag == 'li':
                self.blocks.append(('bullet', text))
            elif self.current_tag in ['h1', 'h2', 'h3', 'h4', 'blockquote', 'pre']:
                tag_map = {'blockquote': 'quote', 'pre': 'code'}
                self.blocks.append((tag_map.get(self.current_tag, self.current_tag), text))
            elif self.current_tag == 'p':
                self.blocks.append(('p', text))

    parser = SimpleBlockParser()
    try:
        parser.feed(content_html)
        parser.flush_buffer()
        if parser.blocks:
            return parser.blocks
    except Exception:
        pass

    return [('p', l.strip()) for l in (content_plain or "").splitlines() if l.strip()]


def generate_note_pdf(note: Note) -> bytes:
    """
    Gera um documento PDF profissional e bem formatado a partir de uma anotação.
    """
    import html
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=54,
        leftMargin=54,
        topMargin=54,
        bottomMargin=54,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'PDFNoteTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=22,
        leading=26,
        textColor=colors.HexColor('#0f172a'),
        spaceAfter=8,
    )

    meta_style = ParagraphStyle(
        'PDFNoteMeta',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor('#64748b'),
        spaceAfter=12,
    )

    body_style = ParagraphStyle(
        'PDFNoteBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10.5,
        leading=16,
        textColor=colors.HexColor('#1e293b'),
        spaceAfter=8,
    )

    h1_style = ParagraphStyle('PDFNoteH1', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=16, leading=20, textColor=colors.HexColor('#0f172a'), spaceBefore=12, spaceAfter=6)
    h2_style = ParagraphStyle('PDFNoteH2', parent=styles['Heading2'], fontName='Helvetica-Bold', fontSize=14, leading=18, textColor=colors.HexColor('#1e293b'), spaceBefore=10, spaceAfter=4)
    h3_style = ParagraphStyle('PDFNoteH3', parent=styles['Heading3'], fontName='Helvetica-Bold', fontSize=12, leading=16, textColor=colors.HexColor('#334155'), spaceBefore=8, spaceAfter=4)

    quote_style = ParagraphStyle(
        'PDFNoteQuote',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=10,
        leading=15,
        textColor=colors.HexColor('#475569'),
        leftIndent=18,
        spaceBefore=6,
        spaceAfter=6,
    )

    code_style = ParagraphStyle(
        'PDFNoteCode',
        parent=styles['Normal'],
        fontName='Courier',
        fontSize=9.5,
        leading=14,
        textColor=colors.HexColor('#0f172a'),
        backColor=colors.HexColor('#f1f5f9'),
        borderPadding=6,
        spaceBefore=6,
        spaceAfter=6,
    )

    bullet_style = ParagraphStyle(
        'PDFNoteBullet',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10.5,
        leading=15,
        textColor=colors.HexColor('#1e293b'),
        leftIndent=16,
        spaceAfter=4,
    )

    story = []

    # Título da Nota
    title = note.title.strip() if note.title else "Sem título"
    story.append(Paragraph(html.escape(title), title_style))

    # Metadados
    nb_title = note.notebook.title if note.notebook else "Sem Caderno"
    created = note.created_at.strftime("%d/%m/%Y %H:%M")
    updated = note.updated_at.strftime("%d/%m/%Y %H:%M")
    tags_list = [t.name for t in note.tags.all()]
    tags_str = f" &bull; Tags: #{', #'.join(tags_list)}" if tags_list else ""
    meta_text = f"Caderno: <b>{html.escape(nb_title)}</b> &bull; Criado: {created} &bull; Atualizado: {updated}{tags_str}"
    story.append(Paragraph(meta_text, meta_style))

    # Linha divisória sutil
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#e2e8f0'), spaceAfter=14))

    # Conteúdo estruturado
    blocks = parse_html_to_blocks(note.content_html, note.content_plain)
    for b_type, b_val in blocks:
        if b_type == 'h1':
            story.append(Paragraph(html.escape(b_val), h1_style))
        elif b_type == 'h2':
            story.append(Paragraph(html.escape(b_val), h2_style))
        elif b_type in ('h3', 'h4'):
            story.append(Paragraph(html.escape(b_val), h3_style))
        elif b_type == 'quote':
            story.append(Paragraph(f"&ldquo;{html.escape(b_val)}&rdquo;", quote_style))
        elif b_type == 'code':
            formatted_code = html.escape(b_val).replace('\n', '<br/>')
            story.append(Paragraph(formatted_code, code_style))
        elif b_type == 'todo':
            checked, text = b_val
            mark = "[✓]" if checked else "[  ]"
            story.append(Paragraph(f"<b>{mark}</b> {html.escape(text)}", bullet_style))
        elif b_type == 'bullet':
            story.append(Paragraph(f"&bull; {html.escape(b_val)}", bullet_style))
        elif b_type == 'hr':
            story.append(HRFlowable(width="100%", thickness=0.8, color=colors.HexColor('#cbd5e1'), spaceBefore=8, spaceAfter=8))
        else:
            if b_val.strip():
                story.append(Paragraph(html.escape(b_val), body_style))
            else:
                story.append(Spacer(1, 6))

    doc.build(story)
    return buffer.getvalue()


def generate_note_docx(note: Note) -> bytes:
    """
    Gera um documento Microsoft Word (.docx) formatado e estilizado a partir de uma anotação.
    """
    import docx
    from docx.shared import Inches, Pt, RGBColor

    doc = docx.Document()

    # Margens da página
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Título da Nota
    title = note.title.strip() if note.title else "Sem título"
    doc.add_heading(title, level=0)

    # Metadados
    nb_title = note.notebook.title if note.notebook else "Sem Caderno"
    created = note.created_at.strftime("%d/%m/%Y %H:%M")
    updated = note.updated_at.strftime("%d/%m/%Y %H:%M")
    tags_list = [t.name for t in note.tags.all()]
    tags_str = f" | Tags: #{', #'.join(tags_list)}" if tags_list else ""

    meta_p = doc.add_paragraph()
    meta_run = meta_p.add_run(f"Caderno: {nb_title} | Criado: {created} | Atualizado: {updated}{tags_str} | skdnote")
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = RGBColor(100, 116, 139)
    meta_p.paragraph_format.space_after = Pt(14)

    # Conteúdo estruturado
    blocks = parse_html_to_blocks(note.content_html, note.content_plain)
    for b_type, b_val in blocks:
        if b_type == 'h1':
            doc.add_heading(b_val, level=1)
        elif b_type == 'h2':
            doc.add_heading(b_val, level=2)
        elif b_type in ('h3', 'h4'):
            doc.add_heading(b_val, level=3)
        elif b_type == 'quote':
            p = doc.add_paragraph(style='Quote')
            p.add_run(b_val)
        elif b_type == 'code':
            p = doc.add_paragraph()
            r = p.add_run(b_val)
            r.font.name = 'Courier New'
            r.font.size = Pt(9.5)
            p.paragraph_format.left_indent = Inches(0.25)
        elif b_type == 'todo':
            checked, text = b_val
            mark = "☑" if checked else "☐"
            doc.add_paragraph(f"{mark} {text}", style='List Bullet')
        elif b_type == 'bullet':
            doc.add_paragraph(b_val, style='List Bullet')
        elif b_type == 'hr':
            p = doc.add_paragraph()
            r = p.add_run("―" * 40)
            r.font.color.rgb = RGBColor(203, 213, 225)
        else:
            if b_val.strip():
                doc.add_paragraph(b_val)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
